"""
Generate sample legal documents in multiple formats for verification testing.
"""

import os
from pathlib import Path
from typing import Dict, Any


def generate_contract_a() -> str:
    """Generate Contract A content."""
    return """
PURCHASE AGREEMENT

Clause 1: Definitions
Party A shall mean the seller of goods under this agreement.
Party B shall mean the buyer of goods under this agreement.

Clause 2: Payment Terms
Client agrees to pay attorney fees of $10,000 within 30 days of invoice.
Payment shall be made by wire transfer or certified check.

Clause 3: Delivery
Seller agrees to delivery within 15 days after payment received.
Delivery shall be made to the address specified by Party B.

Clause 4: Termination
Either party may terminate this agreement with 14 days written notice.
Termination notice must be sent via certified mail.
"""


def generate_case_law() -> str:
    """Generate Case Law Example content."""
    return """
CASE LAW EXAMPLE
Smith v. Jones Commercial Services (2023)

Section 1: Background
This case involves a breach of contract dispute between commercial parties.
The plaintiff alleged that defendant failed to deliver goods within the agreed timeframe.

Section 2: Decision
The court held that timely delivery is a material obligation under commercial contracts.
Failure to deliver within agreed timeframes constitutes a material breach.

Section 3: Legal Principle
Parties must perform within agreed timeframe unless specified otherwise in the contract.
Courts will enforce delivery terms strictly in commercial transactions.
"""


def generate_statute() -> str:
    """Generate Statute Example content."""
    return """
COMMERCIAL TRANSACTIONS ACT, 2023

Section 10: Payment Obligations
All commercial payments must be made within 45 days unless agreed in writing.
Late payments shall incur interest at the statutory rate.

Section 15: Delivery
Goods must be delivered within 30 days unless contract specifies otherwise.
Sellers have a duty to provide reasonable notice of delivery delays.

Section 20: Termination Rights
Parties may terminate with reasonable notice not less than 7 days.
Termination rights are subject to the terms of the specific agreement.
"""


def create_txt_files(docs_dir: Path) -> None:
    """Create TXT versions of all documents."""
    documents = {
        "Contract_A.txt": generate_contract_a(),
        "Case_Law_Example.txt": generate_case_law(),
        "Statute_Example.txt": generate_statute()
    }
    
    for filename, content in documents.items():
        file_path = docs_dir / filename
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(content.strip())
        print(f"✅ Created {filename}")


def create_docx_files(docs_dir: Path) -> None:
    """Create DOCX versions of all documents."""
    try:
        from docx import Document
    except ImportError:
        print("❌ python-docx not installed. Skipping DOCX generation.")
        return
    
    documents = {
        "Contract_A.docx": generate_contract_a(),
        "Case_Law_Example.docx": generate_case_law(),
        "Statute_Example.docx": generate_statute()
    }
    
    for filename, content in documents.items():
        doc = Document()
        # Split content into paragraphs and add to document
        for paragraph in content.strip().split('\n\n'):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        
        file_path = docs_dir / filename
        doc.save(str(file_path))
        print(f"✅ Created {filename}")


def create_pdf_files(docs_dir: Path) -> None:
    """Create PDF versions of all documents."""
    try:
        from fpdf import FPDF
    except ImportError:
        print("❌ fpdf2 not installed. Skipping PDF generation.")
        return
    
    documents = {
        "Contract_A.pdf": generate_contract_a(),
        "Case_Law_Example.pdf": generate_case_law(),
        "Statute_Example.pdf": generate_statute()
    }
    
    for filename, content in documents.items():
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        # Split content into lines and add to PDF
        for line in content.strip().split('\n'):
            if line.strip():
                # Handle long lines by wrapping
                if len(line) > 80:
                    words = line.split()
                    current_line = ""
                    for word in words:
                        if len(current_line + word) < 80:
                            current_line += word + " "
                        else:
                            if current_line:
                                pdf.cell(0, 10, current_line.strip(), ln=True)
                            current_line = word + " "
                    if current_line:
                        pdf.cell(0, 10, current_line.strip(), ln=True)
                else:
                    pdf.cell(0, 10, line, ln=True)
            else:
                pdf.cell(0, 5, "", ln=True)  # Empty line
        
        file_path = docs_dir / filename
        pdf.output(str(file_path))
        print(f"✅ Created {filename}")


def generate_sample_documents(output_dir: str = None) -> None:
    """Generate all sample documents in multiple formats."""
    if output_dir is None:
        output_dir = Path(__file__).parent / "sample_docs"
    else:
        output_dir = Path(output_dir)
    
    output_dir.mkdir(exist_ok=True)
    
    print("📄 Generating sample legal documents...")
    print(f"📁 Output directory: {output_dir}")
    
    # Generate in all formats
    create_txt_files(output_dir)
    create_docx_files(output_dir)
    create_pdf_files(output_dir)
    
    print(f"\n✅ Generated {len(list(output_dir.glob('*')))} files in {output_dir}")


if __name__ == "__main__":
    import sys
    output_dir = sys.argv[1] if len(sys.argv) > 1 else None
    generate_sample_documents(output_dir)
